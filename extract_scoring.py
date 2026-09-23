import pandas as pd
import openpyxl
import json
import re

# 1. YSQ Scoring
ysq_df = pd.read_excel('/Users/matthias/Downloads/1. Schemavragenlijst.xlsx', sheet_name='Blad3', header=None)

ysq_scoring = {}
for col in range(0, ysq_df.shape[1], 2):
    for r in range(2, 50, 10): # Blocks of schemas
        if r < len(ysq_df) and pd.notna(ysq_df.iloc[r, col]):
            schema_name = str(ysq_df.iloc[r, col]).strip()
            if schema_name and schema_name != "nan":
                items = []
                for i in range(2, 7):
                    item_val = ysq_df.iloc[r+i, col]
                    if pd.notna(item_val):
                        try:
                            items.append(int(item_val))
                        except ValueError:
                            pass
                if items:
                    ysq_scoring[schema_name] = items

# 2. SMI Scoring
smi_wb = openpyxl.load_workbook('/Users/matthias/Downloads/2. Modivragenlijst.xlsx', data_only=False)
ws = smi_wb['invoer']

smi_scoring = {}
# Modes listed starting row 137, col 2
for row in range(137, 160):
    mode_name = ws.cell(row=row, column=2).value
    formula = ws.cell(row=row, column=3).value
    
    if mode_name and formula and isinstance(formula, str) and formula.startswith('=AVERAGE('):
        # Extract cell references like C11, C13, etc.
        cells_str = formula.replace('=AVERAGE(', '').replace(')', '')
        cells = cells_str.split(',')
        
        items = []
        for cell in cells:
            cell = cell.strip()
            if cell.startswith('C'):
                row_num = int(cell[1:])
                # Offset is 7 (Row 11 -> Q 4)
                q_num = row_num - 7
                items.append(q_num)
        
        smi_scoring[mode_name] = items

with open('/Users/matthias/.gemini/antigravity/scratch/schema-therapy-app/src/data/ysq-scoring.json', 'w') as f:
    json.dump(ysq_scoring, f, indent=2)

with open('/Users/matthias/.gemini/antigravity/scratch/schema-therapy-app/src/data/smi-scoring.json', 'w') as f:
    json.dump(smi_scoring, f, indent=2)

print("Scoring logic extracted successfully!")
